#!/usr/bin/env python3
"""
Render test data corpus using Jinja2 templates.

Chainable tool that reads JSON from stdin and writes rendered output to stdout.
Can also read from xlsx/JSON files and write to directory.

Features:
- Reads JSON from stdin (or --input file for xlsx/json)
- Writes rendered output to stdout (or --output dir for files)
- Profile-based configuration (count, output dir, template, field mapping)
- Jinja2 template rendering
- Field name mapping for public/internal versions
- ServicePoints parsing from pipe-delimited format
"""

import json
import sys
from pathlib import Path
from typing import Annotated, Optional

import typer
import yaml
from jinja2 import Environment, FileSystemLoader
from openpyxl import load_workbook

# Paths
SCRIPT_DIR = Path(__file__).parent
PROJECT_ROOT = SCRIPT_DIR.parent.parent
TESTDATA_DIR = PROJECT_ROOT / "dat" / "testdata"
CORPUS_FILE = TESTDATA_DIR / "corpus.xlsx"
PROFILES_DIR = TESTDATA_DIR / "profiles"

# Columns that are metadata, not record data
METADATA_COLUMNS = {"TestCaseId", "TestType", "Description", "ExpectedResult", "ExecutionLog"}

app = typer.Typer(help="Render test data corpus using Jinja2 templates.")


def parse_service_points(pipe_string: str) -> list[dict]:
    """
    Parse pipe-delimited service points.

    Format: Type|PostalCode|City per ServicePoint, separated by ;
    Example: "POSTBOX|12345|Berlin;ADDRESS|54321|Munich"
    """
    if not pipe_string or not pipe_string.strip():
        return []

    service_points = []
    for entry in pipe_string.split(";"):
        entry = entry.strip()
        if not entry:
            continue

        parts = entry.split("|")
        if len(parts) >= 3:
            service_points.append({
                "type": parts[0],
                "postalcode": parts[1],
                "city": parts[2],
            })

    return service_points


def load_profile(profile_name: str, profiles_dir: Path) -> dict:
    """Load a profile from the profiles directory."""
    profile_path = profiles_dir / f"{profile_name}.yaml"

    if not profile_path.exists():
        available = [p.stem for p in profiles_dir.glob("*.yaml")]
        raise FileNotFoundError(
            f"Profile '{profile_name}' not found at {profile_path}\n"
            f"Available profiles: {', '.join(available) if available else 'none'}"
        )

    with open(profile_path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def list_profiles(profiles_dir: Path) -> list[str]:
    """List available profile names."""
    return sorted([p.stem for p in profiles_dir.glob("*.yaml")])


def load_field_mapping(mapping_path: Path | None) -> dict:
    """Load field name mapping from YAML file."""
    if not mapping_path or not mapping_path.exists():
        return {}

    with open(mapping_path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f) or {}


def load_corpus_from_xlsx(corpus_path: Path, sheet_name: str = "Synthetic") -> list[dict]:
    """Load test cases from corpus Excel file."""
    if not corpus_path.exists():
        raise FileNotFoundError(f"Corpus file not found: {corpus_path}")

    wb = load_workbook(corpus_path, read_only=True)

    if sheet_name not in wb.sheetnames:
        available = ", ".join(wb.sheetnames)
        raise ValueError(f"Sheet '{sheet_name}' not found. Available: {available}")

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if len(rows) < 2:
        return []

    headers = [str(h).strip() if h else "" for h in rows[0]]
    test_cases = []

    for row in rows[1:]:
        if not row or not any(row):
            continue

        record = {}
        metadata = {}

        for i, value in enumerate(row):
            if i >= len(headers):
                break

            header = headers[i]
            if not header:
                continue

            # Convert value to string, handle None
            str_value = str(value).strip() if value is not None else ""

            if header in METADATA_COLUMNS:
                metadata[header] = str_value
            else:
                record[header] = str_value

        # Parse ServicePoints from pipe format
        if "ServicePoints" in record:
            record["ServicePoints_Parsed"] = parse_service_points(record["ServicePoints"])

        test_cases.append({
            "metadata": metadata,
            "record": record,
            "servicePoints": record.get("ServicePoints_Parsed", []),
        })

    return test_cases


def load_corpus_from_json(json_path: Path) -> list[dict]:
    """Load test cases from JSON file."""
    if not json_path.exists():
        raise FileNotFoundError(f"JSON file not found: {json_path}")

    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    # Normalize to list
    if isinstance(data, dict):
        data = [data]

    return normalize_json_test_cases(data)


def load_corpus_from_stdin() -> list[dict]:
    """Load test cases from JSON on stdin."""
    if sys.stdin.isatty():
        raise ValueError("No input on stdin. Use --input for file input or pipe JSON.")

    data = json.load(sys.stdin)

    # Normalize to list
    if isinstance(data, dict):
        data = [data]

    return normalize_json_test_cases(data)


def normalize_json_test_cases(data: list[dict]) -> list[dict]:
    """Normalize JSON test cases to internal format."""
    test_cases = []

    for item in data:
        # Check if already in {metadata, record, servicePoints} format
        if "metadata" in item and "record" in item:
            tc = {
                "metadata": item.get("metadata", {}),
                "record": item.get("record", {}),
                "servicePoints": item.get("servicePoints", []),
            }
            # Also set ServicePoints_Parsed for template compatibility
            tc["record"]["ServicePoints_Parsed"] = tc["servicePoints"]
        else:
            # Flat format - split into metadata/record
            metadata = {}
            record = {}
            for key, value in item.items():
                if key in METADATA_COLUMNS:
                    metadata[key] = value
                elif key == "servicePoints":
                    continue
                else:
                    record[key] = value

            service_points = item.get("servicePoints", [])
            if not service_points and "ServicePoints" in record:
                service_points = parse_service_points(record["ServicePoints"])

            tc = {
                "metadata": metadata,
                "record": record,
                "servicePoints": service_points,
            }
            tc["record"]["ServicePoints_Parsed"] = service_points

        test_cases.append(tc)

    return test_cases


def load_corpus(input_path: Optional[Path], sheet_name: str = "Synthetic") -> list[dict]:
    """Load test cases from file or stdin."""
    if input_path is None:
        return load_corpus_from_stdin()

    if str(input_path).endswith(".xlsx"):
        return load_corpus_from_xlsx(input_path, sheet_name)
    elif str(input_path).endswith(".json"):
        return load_corpus_from_json(input_path)
    else:
        # Try JSON
        return load_corpus_from_json(input_path)


def parse_row_range(rows_arg: str, total_count: int) -> list[int]:
    """
    Parse row range argument.

    Examples:
        "1" -> [0]
        "1-5" -> [0, 1, 2, 3, 4]
        "all" -> [0, 1, 2, ..., total_count-1]
    """
    if rows_arg.lower() == "all":
        return list(range(total_count))

    if "-" in rows_arg:
        parts = rows_arg.split("-")
        start = int(parts[0]) - 1  # 1-indexed to 0-indexed
        end = int(parts[1])  # inclusive end
        return list(range(max(0, start), min(end, total_count)))

    # Single row
    idx = int(rows_arg) - 1
    if 0 <= idx < total_count:
        return [idx]
    return []


def render_test_case(
    template,
    test_case: dict,
    field_map: dict,
) -> str:
    """Render a single test case using the template."""
    # Support both servicePoints at top level and ServicePoints_Parsed in record
    service_points = test_case.get("servicePoints", [])
    if not service_points:
        service_points = test_case["record"].get("ServicePoints_Parsed", [])

    return template.render(
        record=test_case["record"],
        metadata=test_case["metadata"],
        field_map=field_map,
        service_points=service_points,
    )


@app.command("list")
def list_cmd(
    profiles_dir: Annotated[Optional[Path], typer.Option("--profiles-dir", "-d", help="Profiles directory")] = None,
):
    """List available profiles."""
    pdir = profiles_dir or PROFILES_DIR
    profiles = list_profiles(pdir)

    if not profiles:
        typer.echo(f"No profiles found in {pdir}")
        raise typer.Exit(1)

    typer.echo(f"Available profiles in {pdir}:")
    for profile in profiles:
        typer.echo(f"  - {profile}")


@app.command("render")
def render(
    profile: Annotated[Optional[str], typer.Option("--profile", "-p", help="Profile name (without .yaml extension)")] = None,
    template: Annotated[Optional[Path], typer.Option("--template", "-t", help="Template path (overrides profile)")] = None,
    rows: Annotated[Optional[str], typer.Option("--rows", "-r", help="Row range (e.g., '1', '1-5', 'all')")] = None,
    sheet: Annotated[str, typer.Option("--sheet", "-s", help="Sheet name for xlsx input")] = "Synthetic",
    input_path: Annotated[Optional[Path], typer.Option("--input", "-i", help="Input file (.xlsx, .json). Omit to read JSON from stdin.")] = None,
    output_path: Annotated[Optional[Path], typer.Option("--output", "-o", help="Output directory for files. Omit to write to stdout.")] = None,
    profiles_dir: Annotated[Optional[Path], typer.Option("--profiles-dir", "-d", help="Profiles directory")] = None,
    dry_run: Annotated[bool, typer.Option("--dry-run", help="Show what would be rendered without writing")] = False,
    quiet: Annotated[bool, typer.Option("--quiet", "-q", help="Suppress status messages (for piping)")] = False,
):
    """Render test data using Jinja2 templates.

    By default, reads JSON from stdin and writes rendered output to stdout.
    Use --input for file input (.xlsx or .json).
    Use --output to write files to a directory.
    """
    pdir = profiles_dir or PROFILES_DIR

    def log(msg: str):
        if not quiet:
            print(msg, file=sys.stderr)

    # Must have either profile or template
    if not profile and not template:
        typer.echo("Error: Must specify either --profile or --template", err=True)
        raise typer.Exit(1)

    # Load profile if specified
    if profile:
        log(f"Loading profile: {profile}")
        try:
            profile_config = load_profile(profile, pdir)
        except FileNotFoundError as e:
            typer.echo(f"Error: {e}", err=True)
            raise typer.Exit(1)
        log(f"  Description: {profile_config.get('description', 'N/A')}")
    else:
        profile_config = {}

    # Determine input source
    # Priority: CLI --input > profile input > stdin
    effective_input = input_path
    if effective_input is None and profile_config.get("input"):
        effective_input = TESTDATA_DIR / profile_config["input"]

    # Load corpus from input source
    if effective_input:
        log(f"Loading from: {effective_input}")
    else:
        log("Reading JSON from stdin...")

    try:
        test_cases = load_corpus(effective_input, sheet)
    except (FileNotFoundError, ValueError) as e:
        typer.echo(f"Error: {e}", err=True)
        raise typer.Exit(1)

    log(f"  Found {len(test_cases)} test cases")

    if not test_cases:
        typer.echo("No test cases found.", err=True)
        raise typer.Exit(1)

    # Determine which rows to render
    if rows:
        row_indices = parse_row_range(rows, len(test_cases))
    else:
        count = profile_config.get("count", "all")
        if count == "all" or count is None:
            row_indices = list(range(len(test_cases)))
        else:
            row_indices = list(range(min(int(count), len(test_cases))))

    log(f"  Rendering {len(row_indices)} test cases")

    # Load field mapping
    mapping_path = profile_config.get("field_mapping")
    if mapping_path:
        mapping_full_path = TESTDATA_DIR / mapping_path
        log(f"Loading field mapping from: {mapping_full_path}")
        field_map = load_field_mapping(mapping_full_path)
    else:
        field_map = {}

    # Setup Jinja2 environment
    # Template from CLI overrides profile
    if template:
        template_path = template
        if template_path.is_absolute():
            template_dir = template_path.parent
            template_name = template_path.name
        else:
            # Relative to TESTDATA_DIR
            template_dir = TESTDATA_DIR / template_path.parent
            template_name = template_path.name
    else:
        template_path_str = profile_config.get("template", "templates/internal.j2")
        template_dir = TESTDATA_DIR / Path(template_path_str).parent
        template_name = Path(template_path_str).name
        template_path = template_dir / template_name

    log(f"Loading template: {template_path}")

    if not (template_dir / template_name).exists():
        typer.echo(f"Error: Template not found: {template_dir / template_name}", err=True)
        raise typer.Exit(1)

    env = Environment(
        loader=FileSystemLoader(template_dir),
        trim_blocks=True,
        lstrip_blocks=True,
    )
    jinja_template = env.get_template(template_name)

    # Determine output mode
    # Priority: CLI --output > profile output_dir > stdout
    final_output_dir = output_path
    if final_output_dir is None and profile_config.get("output_dir"):
        final_output_dir = PROJECT_ROOT / profile_config["output_dir"]

    # Convert relative path to absolute
    if final_output_dir and not final_output_dir.is_absolute():
        final_output_dir = PROJECT_ROOT / final_output_dir

    if dry_run:
        output_target = final_output_dir or "stdout"
        log(f"\n[DRY RUN] Would write to: {output_target}")
        log(f"[DRY RUN] Items to render: {len(row_indices)}")
        log("\nSample outputs:")
        for idx in row_indices[:3]:
            test_case = test_cases[idx]
            test_id = test_case["metadata"].get("TestCaseId", f"TC{idx+1:03d}")
            ext = ".json" if "json" in template_name else ".txt"
            log(f"  {test_id}{ext}: {test_case['metadata'].get('Description', 'N/A')}")
        if len(row_indices) > 3:
            log(f"  ... and {len(row_indices) - 3} more")
        return

    if final_output_dir:
        # Write to directory
        final_output_dir.mkdir(parents=True, exist_ok=True)
        log(f"Output directory: {final_output_dir}")

        for idx in row_indices:
            test_case = test_cases[idx]
            test_id = test_case["metadata"].get("TestCaseId", f"TC{idx+1:03d}")

            rendered = render_test_case(jinja_template, test_case, field_map)

            # Determine output extension from template name
            if template_name.endswith(".j2"):
                if "json" in template_name:
                    ext = ".json"
                else:
                    ext = ".txt"
            else:
                ext = ".txt"

            output_file = final_output_dir / f"{test_id}{ext}"
            with open(output_file, "w", encoding="utf-8") as f:
                f.write(rendered)

            log(f"  Written: {output_file.name}")

        log(f"Done! Rendered {len(row_indices)} files to {final_output_dir}")
    else:
        # Write to stdout
        rendered_outputs = []
        for idx in row_indices:
            test_case = test_cases[idx]
            rendered = render_test_case(jinja_template, test_case, field_map)
            rendered_outputs.append({
                "testCaseId": test_case["metadata"].get("TestCaseId", f"TC{idx+1:03d}"),
                "rendered": rendered,
                "metadata": test_case["metadata"],
            })

        # Output as JSON for piping
        print(json.dumps(rendered_outputs, indent=2, ensure_ascii=False))


@app.callback(invoke_without_command=True)
def main(ctx: typer.Context):
    """Render test data corpus using Jinja2 templates."""
    if ctx.invoked_subcommand is None:
        typer.echo("Use --help to see available commands")
        raise typer.Exit(0)


if __name__ == "__main__":
    app()
